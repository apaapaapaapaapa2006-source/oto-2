#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
works.csv から oto-2-works.xlsx を作る。

Google ドライブへ放り込めばそのまま Google スプレッドシートになる。
シートは ゲーム / 楽器 / エフェクト の3枚 + 使い方。
列は  no ・ 〇/△ ・ 日付 ・ 名前 ・ 簡単な説明 ・ memo ・ ジャンル

  python3 tools/make_xlsx.py            # works.csv -> oto-2-works.xlsx
  python3 tools/make_xlsx.py out.xlsx   # 出力先を指定

【重要】〇/△ と memo は人が手で編集する列なので、
作り直すときは既存ファイルからその2列を読み戻して引き継ぐ。
名前（NAME）を鍵にして突き合わせるので、行を並べ替えても消えない。
"""
import csv
import os
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
CSV_PATH = os.path.join(REPO, 'works.csv')

# シート名 -> works.csv の カテゴリ
SHEETS = [('ゲーム', 'ゲーム'), ('楽器', '楽器'), ('エフェクト', 'エフェクト')]
HEADERS = ['no', '〇/△', '日付', '名前', '簡単な説明', 'memo', 'ジャンル']

# 幅は文字数ではなく「読める幅」で決めてある。説明は折り返す前提。
WIDTHS = [5, 7, 12, 18, 68, 30, 11]

FONT = 'Arial'
INK = '1F2933'
HEAD_BG = '1F3A5C'
HEAD_FG = 'FFFFFF'
EDIT_BG = 'FFF6D8'          # 手で編集する列の下地
RULE = 'C8D2DC'

MARKS = ['〇', '△', '×']


def read_works():
    with open(CSV_PATH, encoding='utf-8') as f:
        rows = list(csv.DictReader(f))
    # 制作日の古い順。no を安定させたいので、この順序が台帳の背骨になる。
    rows.sort(key=lambda r: (r['日付'], r['名前']))
    return rows


def carry_over(path):
    """既存ファイルから 名前 -> (〇/△, memo) を回収する。"""
    kept = {}
    if not path or not os.path.exists(path):
        return kept
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as e:                                  # 壊れていても止めない
        print('  既存ファイルを読めなかったので引き継ぎなし: %s' % e)
        return kept
    for ws in wb.worksheets:
        if ws.title not in [s for s, _ in SHEETS]:
            continue
        head = [c.value for c in ws[1]]
        if head[:len(HEADERS)] != HEADERS:
            continue
        for row in ws.iter_rows(min_row=2, values_only=True):
            name = row[3]
            if not name:
                continue
            kept[str(name)] = (row[1], row[5])
    return kept


def style_sheet(ws, rows, kept):
    thin = Side(style='thin', color=RULE)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(HEADERS)
    for i, cell in enumerate(ws[1]):
        cell.font = Font(name=FONT, bold=True, size=10, color=HEAD_FG)
        cell.fill = PatternFill('solid', fgColor=HEAD_BG)
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
    ws.row_dimensions[1].height = 26
    ws['B1'].comment = Comment(
        '手で編集する列。\n〇 = このままでよい\n△ = 後日なおす\n× = 作り直す\n'
        '列の見出しの▽から並べ替え・絞り込みができる。', 'oto-2', width=260, height=110)
    ws['F1'].comment = Comment(
        '手で編集する列。なおしたい点や気づいたことを書く。\n'
        'ここと 〇/△ は作り直しても引き継がれる。', 'oto-2', width=260, height=90)

    for n, r in enumerate(rows, start=1):
        mark, memo = kept.get(r['名前'], ('〇', None))
        ws.append([n, mark if mark else '〇', r['日付'], r['名前'],
                   r['説明'], memo, r['カテゴリ']])
        i = ws.max_row
        for j in range(1, len(HEADERS) + 1):
            c = ws.cell(row=i, column=j)
            c.font = Font(name=FONT, size=10, color=INK)
            c.border = border
            c.alignment = Alignment(vertical='top',
                                    horizontal='center' if j in (1, 2, 3, 7) else 'left',
                                    wrap_text=(j == 5 or j == 6))
        # 名前からその作品へ飛べるようにする
        link = ws.cell(row=i, column=4)
        if r.get('URL'):
            link.hyperlink = r['URL']
            link.font = Font(name=FONT, size=10, color='1155CC', underline='single')
        # 編集する列に下地を敷いて、どこを触ればよいか一目でわかるようにする
        for j in (2, 6):
            ws.cell(row=i, column=j).fill = PatternFill('solid', fgColor=EDIT_BG)
        ws.row_dimensions[i].height = 30

    for j, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w

    last = ws.max_row
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(HEADERS)), last)

    if last >= 2:
        dv = DataValidation(type='list', formula1='"%s"' % ','.join(MARKS),
                            allow_blank=True, showDropDown=False)
        dv.error = '〇 / △ / × のどれかを選んでください'
        dv.prompt = '〇=このままでよい / △=後日なおす / ×=作り直す'
        ws.add_data_validation(dv)
        dv.add('B2:B%d' % last)
    return last


def build_howto(ws, counts):
    ws.column_dimensions['A'].width = 16
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 58

    def put(ref, text, **kw):
        c = ws[ref]
        c.value = text
        c.font = Font(name=FONT, size=kw.get('size', 10), bold=kw.get('bold', False),
                      color=kw.get('color', INK))
        c.alignment = Alignment(vertical='top', wrap_text=kw.get('wrap', False),
                                horizontal=kw.get('h', 'left'))
        return c

    put('A1', 'Oto II 作品台帳', size=16, bold=True)
    put('A2', 'works.csv から自動生成。3つのシートに、ゲーム / 楽器 / エフェクト が分かれて入っている。',
        color='5A6878')

    put('A4', '編集するのはこの2列だけ', bold=True, size=11)
    put('A5', '〇/△')
    put('B5', '〇 = このままでよい　△ = 後日なおす　× = 作り直す（プルダウンから選ぶ）')
    put('A6', 'memo')
    put('B6', 'なおしたい点や気づいたことを書く。長さは自由。')
    put('A7', 'それ以外')
    put('B7', 'works.csv から作られる列なので、書き換えても次の生成で戻る。', color='5A6878')

    put('A9', '並べ替えと絞り込み', bold=True, size=11)
    put('A10', '見出し行の▽から。△だけを表示したいときは 〇/△ の▽で △ にチェック。')
    put('A11', '見出し行は固定してあるので、下までスクロールしても列名が消えない。')

    put('A13', '集計', bold=True, size=11)
    head = ['', 'ゲーム', '楽器', 'エフェクト', '合計']
    for j, h in enumerate(head, start=1):
        c = put(ws.cell(row=14, column=j).coordinate, h, bold=True, h='center')
        if j > 1:
            c.fill = PatternFill('solid', fgColor='EAF0F5')
    marks_rows = {'〇': 15, '△': 16, '×': 17}
    for mark, row in marks_rows.items():
        put('A%d' % row, mark, bold=True, h='center')
        for j, (sheet, _) in enumerate(SHEETS, start=2):
            n = counts[sheet]
            col = get_column_letter(j)
            ws.cell(row=row, column=j).value = (
                '=COUNTIF(%s!$B$2:$B$%d,"%s")' % (sheet, n + 1, mark))
            ws.cell(row=row, column=j).font = Font(name=FONT, size=10, color=INK)
            ws.cell(row=row, column=j).alignment = Alignment(horizontal='center')
        ws.cell(row=row, column=5).value = '=SUM(B%d:D%d)' % (row, row)
        ws.cell(row=row, column=5).font = Font(name=FONT, size=10, bold=True, color=INK)
        ws.cell(row=row, column=5).alignment = Alignment(horizontal='center')
    put('A18', '作品数', bold=True, h='center')
    for j, (sheet, _) in enumerate(SHEETS, start=2):
        ws.cell(row=18, column=j).value = counts[sheet]
        ws.cell(row=18, column=j).font = Font(name=FONT, size=10, color='5A6878')
        ws.cell(row=18, column=j).alignment = Alignment(horizontal='center')
    ws.cell(row=18, column=5).value = '=SUM(B18:D18)'
    ws.cell(row=18, column=5).font = Font(name=FONT, size=10, bold=True, color='5A6878')
    ws.cell(row=18, column=5).alignment = Alignment(horizontal='center')

    put('A20', '作り直しかた', bold=True, size=11)
    put('A21', 'リポジトリで  python3 tools/make_xlsx.py  を実行すると works.csv から作り直す。')
    put('A22', 'そのとき、いまあるファイルの 〇/△ と memo は作品名を鍵にして引き継がれる。'
               'ドライブから .xlsx で書き出して、同じ場所に置いてから実行すること。')
    put('A24', '出どころ: works.csv（このリポジトリ）。日付は JST の制作日。',
        color='5A6878', size=9)


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(REPO, 'oto-2-works.xlsx')
    rows = read_works()
    kept = carry_over(out)
    if kept:
        print('  既存ファイルから %d 件の 〇/△・memo を引き継ぎ' % len(kept))

    wb = Workbook()
    wb.remove(wb.active)
    counts = {}
    for title, cat in SHEETS:
        ws = wb.create_sheet(title)
        sub = [r for r in rows if r['カテゴリ'] == cat]
        style_sheet(ws, sub, kept)
        counts[title] = len(sub)
        print('  %-6s %3d 件' % (title, len(sub)))
    build_howto(wb.create_sheet('使い方'), counts)

    wb.save(out)
    print('書き出し: %s  （合計 %d 件）' % (out, sum(counts.values())))


if __name__ == '__main__':
    main()
